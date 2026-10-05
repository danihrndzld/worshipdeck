"""Tests for pptx_deck_builder. Run: `python3 tests/test_builder.py` (or pytest).

Uses only the committed synthetic reference files, never real church material.
"""
import io
import sys
import tempfile
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
from worshipdeck import builder as B  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

REF = REPO / "reference"


def test_parse_hymn_handles_dot_glued_and_lowercase():
    # "1." separator, a verse number glued onto text ("3¡"), and a lowercase
    # verse start ("4 la") -- all three used to be mishandled.
    raw = "\n".join([
        "283", "BRILLA EN EL SITIO", "Key E Flat.",
        "1. Primera linea del verso uno,", "segunda linea del verso uno.",
        "Coro.- Linea del coro,", "otra linea del coro.",
        "2. Primera del dos,", "segunda del dos.",
        "3¡Glued al numero!", "segunda del tres.",
        "4 la fe empieza en minuscula,", "segunda del cuatro.",
    ])
    p = B.parse_hymn(raw)
    assert len(p["verses"]) == 4, p["verses"]
    assert p["verses"][2][0] == "¡Glued al numero!"
    assert p["verses"][3][0] == "la fe empieza en minuscula,"
    assert p["chorus"][0] == "Linea del coro,"


def test_parse_hymn_accepts_coro_comma_typo():
    raw = "132\nTITULO\n1 Linea uno,\nlinea dos.\nCoro,- :::Linea del coro:::\nOtra del coro."
    assert B.parse_hymn(raw)["chorus"] == [":::Linea del coro:::", "Otra del coro."]


def test_hymn_title_from_caps_header():
    raw = "48\nHALLE UN BUEN AMIGO\nKey F.\n1 Halle un buen amigo,"
    assert B.hymn_title(raw) == "Halle un buen amigo"


def test_hymn_title_joins_wrapped_header_and_keeps_divine_names():
    raw = "7\nCANTAD A CRISTO LOS DE LIMPIO\nCORAZON\nKey G.\n444\n1 Primera linea,"
    assert B.hymn_title(raw) == "Cantad a Cristo los de limpio corazon"


def test_auto_chunk_size_by_line_length():
    assert B.auto_chunk_size(["Linea corta,", "otra corta"]) == 4
    assert B.auto_chunk_size(["Una linea bastante larga para una sola vez,", "x"]) == 2
    parsed = {"verses": [["a,", "b,", "c,", "d,", "e,", "f,", "g,", "h."]], "chorus": ["coro,", "fin."]}
    assert [s[2] for s in B.hymn_sections(parsed, "auto", "auto")] == [4, 4]
    assert [s[2] for s in B.hymn_sections(parsed, 2, "whole")] == [2, "whole"]


def test_capitalize_and_break_long_line():
    assert B.capitalize_first("no hay otro manantial") == "No hay otro manantial"
    assert B.break_long_line("Renuevame, Senor Jesus, pon en mi corazon", 20) == \
        ["Renuevame,", "Senor Jesus,", "pon en mi corazon"]
    assert B.break_long_line("linea corta", 30) == ["linea corta"]
    # No comma: split before a connector word near the middle.
    assert B.break_long_line("Nada me falta pues todo provees", 28) == \
        ["Nada me falta", "pues todo provees"]
    # Never leave a stub piece ("Oh," / "a ti").
    assert B.break_long_line("Oh, tu diestra me ha sostenido", 28) == \
        ["Oh, tu diestra me ha sostenido"]
    assert B.break_long_line("De madrugada yo me acercare a ti", 28) == \
        ["De madrugada yo me acercare a ti"]


def test_apply_format_lowercase_continuation_and_doubled_line():
    assert B.apply_format(["mi alma te anhela y tiene sed"], None) == \
        ["Mi alma te anhela", "y tiene sed"]
    assert B.apply_format(["Oh, tu fidelidad, oh, tu fidelidad"], None) == \
        ["//Oh, tu fidelidad//"]


def test_estimate_lines_wraps_by_width():
    size = 1257554  # 99pt in EMU
    wide = 16666150
    assert B.estimate_lines(["Hermoso eres, mi Senor"], wide, size) == 1
    assert B.estimate_lines(["Levantemos nuestras manos y adoremos"], wide, size) == 2


def test_set_textbox_lines_never_drops_a_line():
    # Build a text box whose 3rd paragraph has NO run (blank separator), then
    # set 3 real lines: the old code dropped the line on the run-less paragraph.
    prs = Presentation(str(REF / "example-template.pptx"))
    box = B.get_textboxes(list(prs.slides)[9])[0]  # scripture text box "1..2"
    tf = box.text_frame
    # ensure there is a run-less paragraph in the middle
    from pptx.oxml.ns import qn as _qn
    blank = tf.paragraphs[0]._p.makeelement(_qn("a:p"), {})
    tf._txBody.insert(list(tf._txBody).index(tf.paragraphs[1]._p), blank)
    B.set_textbox_lines(box, ["uno", "dos", "tres"])
    assert [p.text for p in tf.paragraphs] == ["uno", "dos", "tres"]


def test_clone_slide_copies_background():
    prs = Presentation(str(REF / "example-template.pptx"))
    src = list(prs.slides)[0]
    # give the source slide a <p:bg>
    cSld = src._element.find(qn("p:cSld"))
    from lxml import etree
    bg = etree.SubElement(cSld, qn("p:bg"))
    bgpr = etree.SubElement(bg, qn("p:bgPr"))
    fill = etree.SubElement(bgpr, qn("a:solidFill"))
    clr = etree.SubElement(fill, qn("a:srgbClr")); clr.set("val", "000343")
    etree.SubElement(bgpr, qn("a:effectLst"))
    cSld.insert(0, bg)
    out = Presentation(str(REF / "example-template.pptx"))
    B.delete_all_slides(out)
    new = B.clone_slide(src, out)
    assert new._element.find(qn("p:cSld")).find(qn("p:bg")) is not None


def test_example_build_and_import_deck_roundtrip():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "out.pptx"
        spec = Path(d) / "spec.json"
        spec.write_text(f'''{{
          "template": "{REF}/example-template.pptx",
          "hymnal": "{REF}/example-hymnal.pdf",
          "library": "{REF}/song-library.example.json",
          "output": "{out}",
          "items": [
            {{"op":"clone_range","start":1,"end":3}},
            {{"op":"hymn","himno":1}},
            {{"op":"song","title_white":"Titulo","title_cream":"Cancion",
              "sections":[{{"type":"verse","lines":["Linea uno,","linea dos."]}}]}},
            {{"op":"scripture","book":"Libro","range":"1:1-2",
              "chunks":[["1","Texto uno.","","2","Texto dos."]]}}
          ]
        }}''')
        B.build(str(spec))
        assert zipfile.ZipFile(out).testzip() is None
        Presentation(str(out))  # reopens cleanly
        songs = B.songs_from_deck(str(out))
        titles = [s["title"] for s in songs]
        assert "Titulo Cancion" in titles  # the contemporary song was extracted
        assert not any("Himno" in t for t in titles)  # hymn/scripture skipped
        # Lyric boxes are widened to the frame and centered vertically.
        prs = Presentation(str(out))
        for slide in prs.slides:
            boxes = B.get_textboxes(slide)
            if len(boxes) == 1 and boxes[0].text_frame.text.startswith("Linea uno"):
                box = boxes[0]
                left, _t, width, _h = B.frame_bounds(slide, box)
                assert (box.left, box.width) == (left, width)
                assert abs((box.top + box.height / 2) - prs.slide_height / 2) < 2


def test_sermon_boxes_map_by_layout_not_shape_order():
    # The sermon template is found by layout (lead above the biggest-font
    # title, reference below it), so a logo box further down is ignored and
    # shape order doesn't matter.
    prs = Presentation(str(REF / "example-template.pptx"))
    slide = B.detect_sermon_title_slide(prs)
    assert slide is not None, "sermon-title template not detected"
    lead, title, ref = B.sermon_boxes(B.get_textboxes(slide))
    assert B.first_font_size(title) > B.first_font_size(lead)
    assert B.first_font_size(title) > B.first_font_size(ref)
    assert (lead.top or 0) < (title.top or 0) < (ref.top or 0)
    # The purpose slide has a similar stack but no digits in its bottom box,
    # so it must not win the detection.
    assert "0" in ref.text_frame.text


def test_sermon_op_fills_lead_title_and_reference():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "out.pptx"
        spec = Path(d) / "spec.json"
        spec.write_text(f'''{{
          "template": "{REF}/example-template.pptx",
          "output": "{out}",
          "items": [
            {{"op":"sermon","lead":"Frase de","title":"Ejemplo",
              "reference":"Libro 1"}}
          ]
        }}''')
        B.build(str(spec))
        assert zipfile.ZipFile(out).testzip() is None
        slide = list(Presentation(str(out)).slides)[0]
        lead, title, ref = B.sermon_boxes(B.get_textboxes(slide))
        assert lead.text_frame.text.strip() == "Frase de"
        assert title.text_frame.text.strip() == "Ejemplo"
        assert ref.text_frame.text.strip() == "Libro 1"


def test_import_deck_skips_hymn_titles_in_either_box():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "out.pptx"
        spec = Path(d) / "spec.json"
        spec.write_text(f'''{{
          "template": "{REF}/example-template.pptx",
          "output": "{out}",
          "items": [
            {{"op":"song","title_white":"HIMNO 138","title_cream":"firmes y adelante",
              "sections":[{{"type":"verse","lines":["Linea uno."]}}]}}
          ]
        }}''')
        B.build(str(spec))
        assert B.songs_from_deck(str(out)) == []


def test_insert_adds_scripture_without_touching_existing_slides():
    from lxml import etree
    with tempfile.TemporaryDirectory() as d:
        deck = Path(d) / "deck.pptx"
        out = Path(d) / "out.pptx"
        spec = Path(d) / "spec.json"
        spec.write_text(f'''{{
          "template": "{REF}/example-template.pptx",
          "output": "{deck}",
          "items": [{{"op":"clone_range","start":1,"end":10}}]
        }}''')
        B.build(str(spec))
        spec.write_text(f'''{{
          "deck": "{deck}", "output": "{out}",
          "items": [
            {{"op":"scripture","before":4,"book":"Nuevo","range":"2:3",
              "chunks":[["3","Texto nuevo."]]}},
            {{"op":"scripture","book":"Final","range":"4:5","chunks":[["5","Al final."]]}}
          ]
        }}''')
        B.insert_into_deck(str(spec))
        before = list(Presentation(str(deck)).slides)
        after = list(Presentation(str(out)).slides)
        assert len(after) == len(before) + 4
        texts = [" ".join(b.text_frame.text for b in B.get_textboxes(s)) for s in after]
        assert "Nuevo" in texts[3] and "Texto nuevo." in texts[4]
        assert "Final" in texts[-2] and "Al final." in texts[-1]
        by_id = {s.slide_id: s for s in after}
        for s in before:  # every original slide is byte-identical
            assert etree.tostring(s._element) == etree.tostring(by_id[s.slide_id]._element)


def _run():
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for fn in fns:
        fn()
        print(f"ok  {fn.__name__}")
    print(f"\n{len(fns)} passed")


if __name__ == "__main__":
    _run()
