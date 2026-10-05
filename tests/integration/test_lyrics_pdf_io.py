"""Lyrics PDF tests that build a small deck and render the PDF on disk.
Never uses real church material.
"""
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pypdf import PdfReader

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from worshipdeck import lyrics_pdf as L  # noqa: E402


def add_box(slide, text, size, top):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(9), Inches(1))
    lines = text.split("\n")
    box.text_frame.text = lines[0]
    for line in lines[1:]:
        box.text_frame.add_paragraph().text = line
    for p in box.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)


def make_deck(path, songs):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for title, slides in songs:
        s = prs.slides.add_slide(blank)
        add_box(s, title[0], 116, 1)
        add_box(s, title[1], 240, 3)
        for text in slides:
            add_box(prs.slides.add_slide(blank), text, 99, 2)
    # a sermon passage at the end must not be read as lyrics
    s = prs.slides.add_slide(blank)
    add_box(s, "Libro", 230, 1)
    add_box(s, "4:18-27", 165, 3)
    add_box(prs.slides.add_slide(blank), "18\nTexto del versiculo.", 62, 2)
    prs.save(path)


def test_ep_62_read_deck_finds_songs_and_skips_scripture():
    with tempfile.TemporaryDirectory() as d:
        deck = Path(d) / "DIA 1.pptx"
        make_deck(deck, [(("Cancion", "uno"), ["Linea a\nLinea b", "//Linea c//"])])
        songs = L.read_deck(deck)
        assert [s["title"] for s in songs] == ["Cancion uno"]
        assert songs[0]["slides"] == [["Linea a", "Linea b"], ["//Linea c//"]]


def test_ac_14_one_song_per_page():
    with tempfile.TemporaryDirectory() as d:
        deck, out = Path(d) / "DIA 1.pptx", Path(d) / "out.pdf"
        make_deck(deck, [(("Primera", "cancion"), ["Linea uno\nLinea dos"]),
                         (("Segunda", "cancion"), ["Linea tres"])])
        L.render(L.read_deck(deck), out, "Dia 1", REPO / "reference/example-hymnal.pdf")
        pages = [p.extract_text() for p in PdfReader(out).pages]
        assert len(pages) == 2
        assert "Primera cancion" in pages[0] and "Segunda cancion" in pages[1]


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
