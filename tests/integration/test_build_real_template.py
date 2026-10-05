"""Spec -> PPTX with the church's real template, hymnal.json and song library.

Oracle: the deck the operator actually delivered on 6 September (reference/decks/).
Hymn 64 there is a characterization target: same slides, same lines, same grouping
(first-letter case aside: capitalize_lines is a later church convention).
"""
import json
import zipfile

import pytest
from conftest import REAL_HYMNAL, REAL_TEMPLATE, REPO
from pptx import Presentation

from worshipdeck import builder as B
from worshipdeck import flyer as F

SEPT6 = REPO / "reference" / "decks" / "SEPTIEMBRE 6.pptx"
FLYER_TEXT = "Himno 64\nEres mi amigo fiel\nHimno 83\nAl estar ante Ti\nEl Señor es mi pastor"


def slide_lines(slide):
    return [p.text for box in B.get_textboxes(slide) for p in box.text_frame.paragraphs]


def fold(lines):
    return [line[:1].lower() + line[1:] for line in lines]


@pytest.fixture(scope="module")
def sept6_deck(tmp_path_factory):
    out = tmp_path_factory.mktemp("deck") / "SEPTIEMBRE 6.pptx"
    spec = F.make_spec(FLYER_TEXT, str(out), "Presencia de Dios", "Sed por la", "Salmos 42:1-2 & 63:1-3")
    path = out.with_suffix(".json")
    path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    B.build(str(path))
    return out


def test_cl_12_output_is_a_valid_pptx_whose_images_all_resolve(sept6_deck):
    assert zipfile.ZipFile(sept6_deck).testzip() is None
    prs = Presentation(str(sept6_deck))
    for slide in prs.slides:
        for rel in slide.part.rels.values():
            if "image" in rel.reltype:
                assert rel.target_part.blob, f"broken image on slide {prs.slides.index(slide) + 1}"


def test_ac_04_service_order_follows_the_flyer(sept6_deck):
    texts = [" ".join(slide_lines(s)) for s in Presentation(str(sept6_deck)).slides]
    markers = ["propósito", "Himno 64", "Amigo Fiel", "Himno 83", "Ante ti", "Mi Pastor",
               "Presencia de Dios", "Como el ciervo"]
    where = [next(i for i, t in enumerate(texts) if m in t) for m in markers]
    assert where == sorted(where), dict(zip(markers, where, strict=True))
    assert len(texts) == 75  # 3 intro + 17 + 11 + 19 + 13 + 6 + sermon + 2 + 3 passage slides


def test_ac_05_hymn_64_matches_the_operators_deck(sept6_deck):
    ours = [slide_lines(s) for s in list(Presentation(str(sept6_deck)).slides)[3:20]]
    hers = [slide_lines(s) for s in list(Presentation(str(SEPT6)).slides)[3:20]]
    assert [line.lower() for line in ours[0]] == [line.lower() for line in hers[0]]  # title slide
    assert [fold(s) for s in ours[1:]] == [fold(s) for s in hers[1:]]  # 16 lyric slides


def test_ac_06_psalm_slides_match_the_operators_deck(sept6_deck):
    ours = [slide_lines(s) for s in list(Presentation(str(sept6_deck)).slides)[-5:]]
    hers = [slide_lines(s) for s in list(Presentation(str(SEPT6)).slides)[61:66]]
    assert ours == hers  # Salmos 42:1-2 and 63:1-3: reference slides, verse text, grouping


def test_cl_13_every_hymn_line_reaches_a_slide_once_per_sung_stanza(sept6_deck):
    parsed, _ = B.load_hymn(REAL_HYMNAL, 83)
    sung = [line for _, lines, _ in B.hymn_sections(parsed, "auto", "auto") for line in lines]
    slides = list(Presentation(str(sept6_deck)).slides)
    start = next(i for i, s in enumerate(slides) if "Himno 83" in slide_lines(s)) + 1
    shown = [line for s in slides[start:start + 40] for line in slide_lines(s)][:len(sung)]
    assert fold(shown) == fold(sung)


def test_cl_14_the_deck_keeps_the_templates_slide_size(sept6_deck):
    built, template = Presentation(str(sept6_deck)), Presentation(str(REAL_TEMPLATE))
    assert (built.slide_width, built.slide_height) == (template.slide_width, template.slide_height)
