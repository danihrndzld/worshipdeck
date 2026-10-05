"""Builder tests that write real .pptx files from the synthetic reference template
(build, import-deck, insert, background cloning).
"""
import sys
import tempfile
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

from worshipdeck import builder as B  # noqa: E402

REF = REPO / "reference"


def _run():
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for fn in fns:
        fn()
        print(f"ok  {fn.__name__}")
    print(f"\n{len(fns)} passed")


def test_clone_slide_copies_background():
    prs = Presentation(str(REF / "example-template.pptx"))
    src = next(iter(prs.slides))
    # give the source slide a <p:bg>
    cSld = src._element.find(qn("p:cSld"))
    from lxml import etree
    bg = etree.SubElement(cSld, qn("p:bg"))
    bgpr = etree.SubElement(bg, qn("p:bgPr"))
    fill = etree.SubElement(bgpr, qn("a:solidFill"))
    clr = etree.SubElement(fill, qn("a:srgbClr"))
    clr.set("val", "000343")
    etree.SubElement(bgpr, qn("a:effectLst"))
    cSld.insert(0, bg)
    out = Presentation(str(REF / "example-template.pptx"))
    B.delete_all_slides(out)
    new = B.clone_slide(src, out)
    assert new._element.find(qn("p:cSld")).find(qn("p:bg")) is not None


def test_example_build_and_import_deck_roundtrip():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "out.pptx"
        spec = Path(d) / "spec.json"
        spec.write_text(f'''{{
          "template": "{REF}/example-template.pptx",
          "hymnal": "{REF}/example-hymnal.pdf",
          "library": "{REF}/song-library.example.json",
          "output": "{out}",
          "items": [
            {{"op":"clone_range","start":1,"end":3}},
            {{"op":"hymn","himno":1}},
            {{"op":"song","title_white":"Titulo","title_cream":"Cancion",
              "sections":[{{"type":"verse","lines":["Linea uno,","linea dos."]}}]}},
            {{"op":"scripture","book":"Libro","range":"1:1-2",
              "chunks":[["1","Texto uno.","","2","Texto dos."]]}}
          ]
        }}''')
        B.build(str(spec))
        assert zipfile.ZipFile(out).testzip() is None
        Presentation(str(out))  # reopens cleanly
        songs = B.songs_from_deck(str(out))
        titles = [s["title"] for s in songs]
        assert "Titulo Cancion" in titles  # the contemporary song was extracted
        assert not any("Himno" in t for t in titles)  # hymn/scripture skipped
        # Lyric boxes are widened to the frame and centered vertically.
        prs = Presentation(str(out))
        for slide in prs.slides:
            boxes = B.get_textboxes(slide)
            if len(boxes) == 1 and boxes[0].text_frame.text.startswith("Linea uno"):
                box = boxes[0]
                left, _t, width, _h = B.frame_bounds(slide, box)
                assert (box.left, box.width) == (left, width)
                assert abs((box.top + box.height / 2) - prs.slide_height / 2) < 2


def test_sermon_op_fills_lead_title_and_reference():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "out.pptx"
        spec = Path(d) / "spec.json"
        spec.write_text(f'''{{
          "template": "{REF}/example-template.pptx",
          "output": "{out}",
          "items": [
            {{"op":"sermon","lead":"Frase de","title":"Ejemplo",
              "reference":"Libro 1"}}
          ]
        }}''')
        B.build(str(spec))
        assert zipfile.ZipFile(out).testzip() is None
        slide = next(iter(Presentation(str(out)).slides))
        lead, title, ref = B.sermon_boxes(B.get_textboxes(slide))
        assert lead.text_frame.text.strip() == "Frase de"
        assert title.text_frame.text.strip() == "Ejemplo"
        assert ref.text_frame.text.strip() == "Libro 1"


def test_import_deck_skips_hymn_titles_in_either_box():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "out.pptx"
        spec = Path(d) / "spec.json"
        spec.write_text(f'''{{
          "template": "{REF}/example-template.pptx",
          "output": "{out}",
          "items": [
            {{"op":"song","title_white":"HIMNO 138","title_cream":"firmes y adelante",
              "sections":[{{"type":"verse","lines":["Linea uno."]}}]}}
          ]
        }}''')
        B.build(str(spec))
        assert B.songs_from_deck(str(out)) == []


def test_insert_adds_scripture_without_touching_existing_slides():
    from lxml import etree
    with tempfile.TemporaryDirectory() as d:
        deck = Path(d) / "deck.pptx"
        out = Path(d) / "out.pptx"
        spec = Path(d) / "spec.json"
        spec.write_text(f'''{{
          "template": "{REF}/example-template.pptx",
          "output": "{deck}",
          "items": [{{"op":"clone_range","start":1,"end":10}}]
        }}''')
        B.build(str(spec))
        spec.write_text(f'''{{
          "deck": "{deck}", "output": "{out}",
          "items": [
            {{"op":"scripture","before":4,"book":"Nuevo","range":"2:3",
              "chunks":[["3","Texto nuevo."]]}},
            {{"op":"scripture","book":"Final","range":"4:5","chunks":[["5","Al final."]]}}
          ]
        }}''')
        B.insert_into_deck(str(spec))
        before = list(Presentation(str(deck)).slides)
        after = list(Presentation(str(out)).slides)
        assert len(after) == len(before) + 4
        texts = [" ".join(b.text_frame.text for b in B.get_textboxes(s)) for s in after]
        assert "Nuevo" in texts[3] and "Texto nuevo." in texts[4]
        assert "Final" in texts[-2] and "Al final." in texts[-1]
        by_id = {s.slide_id: s for s in after}
        for s in before:  # every original slide is byte-identical
            assert etree.tostring(s._element) == etree.tostring(by_id[s.slide_id]._element)


if __name__ == "__main__":
    _run()
