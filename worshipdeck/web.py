"""HTTP API + the single-page app in web/.

  uv run flask --app worshipdeck.web run        # http://127.0.0.1:5000

OCR runs in the browser (tesseract.js) so the server needs no tesseract binary
(Vercel has none); the server still finds ink bands OCR skipped, crops doubtful
lines, matches, builds the deck and the lyrics PDF. Corrections are a re-post of
the edited lines to /api/parse: the same re-check the CLI's ask loop does.
"""
import base64
import datetime
import io
import json
import os
import re
import tempfile
from pathlib import Path

from flask import Flask, abort, jsonify, request, send_file, send_from_directory
from pptx import Presentation

from worshipdeck import builder as B
from worshipdeck import flyer as F
from worshipdeck import lyrics_pdf as L

WEB = Path(__file__).resolve().parent.parent / "web"
MONTHS = [m.lower() for m in F.MONTHS]
OPS = {"clone_range", "hymn", "song", "sermon", "scripture"}

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = 12 * 1024 * 1024


# ------------------------------------------------------------------ data access

def hymns():
    if not hasattr(hymns, "cache"):
        hymns.cache = json.loads(Path(B.DEFAULT_HYMNAL).read_text(encoding="utf-8"))["hymns"]
    return hymns.cache


def library():
    return B.load_library(None, B.DEFAULT_LIBRARY_DIR)


def song_title(entry):
    return entry.get("title") or " ".join(filter(None, [entry.get("title_white"), entry.get("title_cream")]))


def sunday(today=None):
    today = today or datetime.date.today()
    return today + datetime.timedelta(days=(6 - today.weekday()) % 7)


def date_label(day):
    return f"Domingo {day.day} de {MONTHS[day.month - 1]}"


def view(item, lib):
    """What the order-of-service screen shows for one spec item."""
    op = item["op"]
    if op == "hymn":
        h = hymns().get(str(item["himno"]), {})
        return {"kind": "himno", "label": f"Himno {item['himno']}", "title": item.get("title") or h.get("title", ""),
                "review": bool(h.get("needs_review"))}
    if op == "song":
        entry = lib.get(item["key"], {}) if "key" in item else item
        return {"kind": "cancion", "label": "Canción", "title": song_title(entry),
                "missing": "key" in item and item["key"] not in lib}
    if op == "scripture":
        todo = any(c == [F.SCRIPTURE_TODO] for c in item.get("chunks") or [])
        return {"kind": "pasaje", "label": "Pasaje", "title": f"{item['book']} {item['range']}", "todo": todo}
    if op == "sermon":
        return {"kind": "predica", "label": "Prédica", "title": item.get("title", "")}
    return {"kind": "intro", "label": "Intro", "title": "Logo y declaración de propósito"}


# ------------------------------------------------------------------ input checks

def image_file(data_url, tmp):
    """Decode the browser's data URL into a file the OCR helpers can open."""
    if not data_url:
        return None
    raw = base64.b64decode(data_url.split(",", 1)[-1])
    if raw[4:12] in (b"ftypheic", b"ftypheix", b"ftypmif1"):
        abort(415, "HEIC no se puede leer aquí: cambia la cámara a JPG o manda una captura")
    path = Path(tmp) / "flyer.img"
    path.write_bytes(raw)
    return path


def clean_spec(spec):
    """Only what a browser may decide: items and a file name. Paths (template,
    hymnal, library, clone_range source) never come from the client."""
    items = []
    for item in spec.get("items", []):
        if item.get("op") not in OPS:
            abort(400, f"op no permitida: {item.get('op')!r}")
        item = {k: v for k, v in item.items() if k != "source"}
        if item["op"] == "song" and "key" in item and item["key"] not in library():
            abort(400, f"la canción {item['key']!r} no está en la biblioteca")
        items.append(item)
    name = re.sub(r"[^\w .áéíóúñÁÉÍÓÚÑ-]", "", Path(spec.get("output") or F.next_sunday_name()).name)
    return {"output": (name if name.endswith(".pptx") else name + ".pptx"), "items": items,
            "format": spec.get("format") or {}}


def reported(fn, *args):
    """The builder and the flyer parser report bad input with SystemExit."""
    try:
        return fn(*args)
    except SystemExit as e:
        abort(400, str(e.code))


def build_to(spec, tmp):
    spec = clean_spec(spec)
    out = Path(tmp) / "deck.pptx"
    path = Path(tmp) / "spec.json"
    path.write_text(json.dumps({**spec, "output": str(out)}, ensure_ascii=False), encoding="utf-8")
    reported(B.build, str(path))
    return spec, out


# ------------------------------------------------------------------ API

@app.post("/api/parse")
def parse():
    body = request.get_json(force=True)
    with tempfile.TemporaryDirectory() as tmp:
        img = image_file(body.get("image"), tmp)
        if body.get("text") is not None:
            lines = body["text"]
        elif body.get("lines") is not None:
            lines = [{**line, "box": tuple(line["box"])} for line in body["lines"]]
            if body.get("fresh") and img:
                lines = F.with_unread_bands(img, lines)
        elif img:
            lines = reported(F.ocr_lines, img)  # local dev with tesseract installed
        else:
            abort(400, "manda image, lines o text")
        if not isinstance(lines, str):
            lines = [{**line, "id": line.get("id", i)} for i, line in enumerate(lines)]
        spec = reported(F.make_spec, lines, body.get("output"), body.get("sermon_title"),
                        body.get("sermon_lead") or "El tema de hoy", body.get("passage"))
        pendientes = []
        for n, p in enumerate(spec.pop("_pendientes", [])):
            if isinstance(p, str):
                p = {"id": f"t{n}", "text": p, "conf": 100, "matched": False}
            elif img:
                buf = Path(tmp) / f"crop{n}.png"
                F.crop_line(img, p["box"], buf)
                p = {**p, "crop": "data:image/png;base64," + base64.b64encode(buf.read_bytes()).decode()}
            pendientes.append(p)
    lib = library()
    day = sunday()
    return jsonify({
        "spec": spec, "views": [view(i, lib) for i in spec["items"]], "pendientes": pendientes,
        "lines": lines if not isinstance(lines, str) else None,
        "warnings": F.warnings({**spec, "_pendientes": []}),
        "date": day.isoformat(), "date_label": date_label(day),
    })


@app.get("/api/search")
def search():
    q = request.args.get("q", "").strip()
    if not q:
        return jsonify([])
    n = F.norm(q)
    out = []
    if q.isdigit() and q in hymns():
        h = hymns()[q]
        out.append({"op": "hymn", "himno": int(q), "label": f"Himno {q}", "title": h["title"], "kind": "himno"})
    for num, h in hymns().items():
        if n and n in F.norm(h["title"]) and len(out) < 12:
            out.append({"op": "hymn", "himno": int(num), "label": f"Himno {num}", "title": h["title"], "kind": "himno"})
    for key, e in library().items():
        if n and (n in F.norm(song_title(e)) or n in key.replace("-", " ")):
            out.insert(0, {"op": "song", "key": key, "label": "Canción", "title": song_title(e), "kind": "cancion"})
    return jsonify(out[:12])


@app.get("/api/hymn/<int:number>")
def hymn(number):
    if str(number) not in hymns():
        abort(404, f"Himno {number} no existe")
    parsed, title = B.load_hymn(B.DEFAULT_HYMNAL, number)
    sections = [{"kind": k, "lines": lines} for k, lines, _ in B.hymn_sections(parsed, 99, 99)]
    return jsonify({"number": number, "title": title, "sections": sections,
                    "review": bool(hymns()[str(number)].get("needs_review"))})


@app.get("/api/song/<key>")
def song(key):
    entry = library().get(key) or abort(404, f"no hay canción {key!r}")
    sections = [{"kind": k, "lines": lines} for k, lines, _ in B.library_sections(entry)]
    return jsonify({"key": key, "title": song_title(entry), "sections": sections})


@app.get("/api/library")
def library_list():
    lib = library()
    return jsonify([{"key": k, "title": song_title(e), "artist": e.get("artist"), "source": e.get("source"),
                     "slides": len(e.get("sections", [])), "date_added": e.get("date_added")}
                    for k, e in sorted(lib.items(), key=lambda kv: song_title(kv[1]).lower())])


@app.post("/api/library")
def library_add():
    if os.environ.get("VERCEL"):
        abort(501, "en el deploy la biblioteca es de solo lectura: la canción va dentro del culto de esta semana")
    entry = request.get_json(force=True)
    if not song_title(entry) or not entry.get("sections"):
        abort(400, "falta el título o la letra")
    dest = B.save_song(entry, overwrite=bool(entry.get("overwrite")))
    if dest is None:
        abort(409, "ya existe una canción con ese título")
    return jsonify({"key": dest.stem})


@app.delete("/api/library/<key>")
def library_delete(key):
    if os.environ.get("VERCEL"):
        abort(501, "en el deploy la biblioteca es de solo lectura")
    path = B.library_dir() / f"{Path(key).name}.json"
    if not path.exists():
        abort(404)
    path.unlink()
    return jsonify({"deleted": key})


@app.post("/api/preview")
def preview():
    with tempfile.TemporaryDirectory() as tmp:
        spec, out = build_to(request.get_json(force=True), tmp)
        slides = []
        for s in Presentation(str(out)).slides:
            lines = [p.text for box in B.get_textboxes(s) for p in box.text_frame.paragraphs if p.text.strip()]
            slides.append({"lines": lines, "pending": F.SCRIPTURE_TODO in lines,
                           "title": len(B.get_textboxes(s)) >= 2 or not lines})
    lib = library()
    views = [view(i, lib) for i in spec["items"]]
    pend = [f"{v['title']}: falta el texto del pasaje" for v in views if v.get("todo")]
    pend += [f"{v['label']} · {v['title']}: el himnario lo marca para revisar" for v in views if v.get("review")]
    return jsonify({"output": spec["output"], "slides": slides, "pendientes": pend})


@app.post("/api/build")
def build():
    with tempfile.TemporaryDirectory() as tmp:
        spec, out = build_to(request.get_json(force=True), tmp)
        data = io.BytesIO(out.read_bytes())
    return send_file(data, as_attachment=True, download_name=spec["output"],
                     mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation")


@app.post("/api/letras")
def letras():
    if request.form.get("reviewed") != "true":
        abort(400, "el PDF de letras solo sale de una presentación ya revisada")
    deck = request.files.get("deck") or abort(400, "falta la presentación")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / Path(deck.filename or "deck.pptx").name
        deck.save(path)
        songs = L.read_deck(path)
        if not songs:
            abort(422, "no encontré canciones en esa presentación")
        out = Path(tmp) / f"Letras - {path.stem}.pdf"
        L.render(songs, out, path.stem.capitalize(), str(B.DEFAULT_HYMNAL))
        data = io.BytesIO(out.read_bytes())
    return send_file(data, as_attachment=True, download_name=out.name, mimetype="application/pdf")


@app.post("/api/import-deck")
def import_deck():
    deck = request.files.get("deck") or abort(400, "falta la presentación")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "deck.pptx"
        deck.save(path)
        songs = B.songs_from_deck(path)
    lib = library()
    for s in songs:
        s["key"] = B.slugify(s["title"])
        s["exists"] = s["key"] in lib
    return jsonify(songs)


@app.errorhandler(400)
@app.errorhandler(404)
@app.errorhandler(409)
@app.errorhandler(415)
@app.errorhandler(422)
@app.errorhandler(501)
def error(e):
    return jsonify({"error": e.description}), e.code


# ------------------------------------------------------------------ the app

@app.get("/")
def index():
    return send_from_directory(WEB, "index.html")


@app.get("/<path:name>")
def static_file(name):
    return send_from_directory(WEB, name)
